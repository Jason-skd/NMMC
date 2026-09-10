"""结果聚合：论文表1/表2、诊断量、result1.xlsx 导出（模板对位填写）。"""
from pathlib import Path

import numpy as np
import openpyxl

from model import ETA, Solution, T

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs" / "附件" / "附件5" / "result1.xlsx"
OUTPUT = ROOT / "output"

PAPER_SLOTS = [6 * h + 1 for h in (10, 12, 14, 16, 18, 20)]  # t = 61..121


def interval_label(t: int) -> str:
    def hm(m):
        return f"{m // 60:02d}:{m % 60:02d}"
    return f"{hm((t - 1) * 10)}-{hm(t * 10)}"


def table1_rows(sol: Solution) -> list[tuple[str, float]]:
    rows = [(interval_label(t), sol.g[t - 1]) for t in PAPER_SLOTS]
    rows += [("全天购电量", sol.g.sum()), ("全天购电费", sol.cost)]
    return rows


def table2_rows(sol: Solution) -> list[tuple[str, float, float]]:
    """6 个 4h 块的 (充电量, 放电量)，电池侧口径（eta*c 与 d/eta）。"""
    rows = []
    for k in range(6):
        seg = slice(24 * k, 24 * (k + 1))
        rows.append((f"{4 * k}:00-{4 * (k + 1)}:00",
                     ETA * sol.c[seg].sum(), sol.d[seg].sum() / ETA))
    return rows


def diagnostics(sol: Solution, p, ell, r) -> dict:
    return {
        "Z_base 无电池直买": float(p @ np.maximum(ell - r, 0.0)),
        "弃电量 Σw": float(sol.w.sum()),
        "逐时段 w ≤ r": bool((sol.w <= r + 1e-9).all()),
        "充放互斥 max(c·d)": float((sol.c * sol.d).max()),
        "全天损耗": float((1 - ETA) * sol.c.sum() + (1 / ETA - 1) * sol.d.sum()),
        "能量平衡残差": float(np.abs(sol.g + r + sol.d - ell - sol.c - sol.w).max()),
        "SOC 范围": (float(sol.S.min()), float(sol.S.max())),
        "循环误差 |S144-S0|": float(abs(sol.S[T] - sol.S[0])),
    }


def write_result1(sol: Solution, path: Path = OUTPUT / "result1.xlsx") -> None:
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb["计划购电量"]
    for i, val in enumerate(sol.g):          # 模板行 i+1 ↔ 时段 t=i+1
        ws.cell(row=2 + i, column=2, value=float(val))
    ws = wb["充放电量"]
    for k, (_, ch_bat, dis_bat) in enumerate(table2_rows(sol)):
        ws.cell(row=2 + k, column=2, value=float(ch_bat))
        ws.cell(row=2 + k, column=3, value=float(dis_bat))
    ws.cell(row=2, column=5, value=float(sol.S[0]))   # 0:00 储电量
    ws.cell(row=3, column=5, value=float(sol.S[T]))   # 24:00 储电量
    path.parent.mkdir(exist_ok=True)
    wb.save(path)
