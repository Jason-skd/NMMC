"""题4 结算与交付：result4-2 / result4-3 双线、C4 自检、价格诊断、摘要。

交付口径（文档 §9）：两线与 result2 / result3 完全同构（行列号对位不改
标签），差别只在费用列一律用**实现价**（C4-4）——计划购电量表「全天购电费」
= 名义计划费 Σp^(j)·g（事后结算量，读法 R 特有口径，A23）；调整购电表
（仅 4-3）= 实际结算购电费 Σ[p·a + ½p|a−g|]（含偏差、不含紧急）。两线
费用列口径不同，分别声明。诊断量对应文档 §12.4；年末可比性修正用逐日
ν^end^(365)（C4-15，不是常数 0.472）。
"""
import datetime as dt
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from scipy.stats import skew

from nmmc.data import ATTACH
from nmmc.data.price import YearPrice
from nmmc.data.year import YearData
from nmmc.model.lp import ETA, T
from nmmc.model.q2 import KAPPA
from nmmc.report.common import (PAPER_SLOTS, md_table, py_date,
                                write_battery_sheet, write_emergency_sheet)
from nmmc.report.q2 import (ATTR_LABELS, EPS_EMERG, PRICE_BANDS, PAPER_DATES,
                            attribute_emergency, report_days)
from nmmc.run.q4 import DayResult

TEMPLATE_42 = ATTACH / "附件5" / "result4-2.xlsx"
TEMPLATE_43 = ATTACH / "附件5" / "result4-3.xlsx"
SOC_EPS = 1e-6        # 对账容差，kWh


def _py_date(day: DayResult, year: YearData) -> dt.date:
    """DayResult 的 0 基日索引转 Python 日期（common.py 的便捷包装）。"""
    return py_date(day.j, year)


def billing_price(price: YearPrice, typ_p: np.ndarray, mode: str,
                  j: int) -> np.ndarray:
    """当日结算价：a1 退化口径用附件1 常数价（与 rollout 的 p_day 一致，
    C4-2 复现要求费列逐位相同），其余口径用附件4 实现价（C4-4）。"""
    return typ_p if mode == "a1" else price.p[j]


def cash_recheck(day: DayResult, price: YearPrice, typ_p: np.ndarray,
                 mode: str) -> float:
    """C4-4 结算价核对：按当日结算价独立重算现金费用与 day.z 的残差。"""
    p = billing_price(price, typ_p, mode, day.j)
    base = p @ day.g if day.a is None else p @ day.a
    if day.a is not None:
        base += 0.5 * (p @ np.abs(day.a - day.g))
    return float(abs(base + KAPPA * (p @ day.u) - day.z))


def nu_band_check(days: list[DayResult], price: YearPrice) -> int:
    """C4-8 ν 区间事后核验：ν^end^(j) ∈ [p_min^(j+1)/η, η·p_max^(j+1)]
    的违例天数（j ≤ 363；末日按 §4.3 边界约定跳过）。"""
    bad = 0
    for day in days:
        if day.j + 1 >= len(price.p):
            continue
        nxt = price.p[day.j + 1]
        if not (nxt.min() / ETA <= day.nu_end <= ETA * nxt.max()):
            bad += 1
    return bad


def day_checks(day: DayResult, year: YearData) -> dict[str, float | int]:
    """单日自检量（文档 §13 的 C4-14 / 对账 / 序关系，继承题2 档 C6–C9）。

    Q4-2（a 为 None）的供给量为 g；Q4-3 为 a（B_0 段前缀 a ≡ g）。
    """
    net = year.ell[day.j] - year.r[day.j]
    supply = day.g if day.a is None else day.a
    d_soc = day.S[T] - day.S[0]
    bat = ETA * day.c.sum() - day.d.sum() / ETA
    losses = (1 - ETA) * day.c.sum() + (1 / ETA - 1) * day.d.sum()
    c7 = abs((supply + day.u - net - day.w).sum() - (d_soc + losses))
    tol = 1e-4
    bad_order = int(day.z_pi > day.z_expost + tol
                    or day.z_expost > day.z + tol
                    or day.z_expost > day.z_ce + tol)
    prefix = 0.0 if day.a is None else float(np.abs(day.a[:36] - day.g[:36]).max())
    return {
        "c6": float(abs(bat - d_soc)),
        "c7": float(c7),
        "prefix": prefix,
        "cd_max": float((day.c * day.d).max()),
        "uw_max": float((day.u * day.w).max()),
        "u_and_c": int(((day.u > EPS_EMERG) & (day.c > EPS_EMERG)).sum()),
        "c9_bad": bad_order,
    }


def year_end_adjust(days: list[DayResult]) -> float:
    """年末可比性修正（文档 §8.2）：Z^adj = ΣZ − ν^end^(365)·(S_144^(365)
    − S_0^(32))，报告期口径；修正系数用逐日 ν（C4-15，非 0.472）。"""
    rep = report_days(days)
    z = sum(d.z for d in rep)
    return float(z - rep[-1].nu_end * (rep[-1].S[T] - rep[0].S[0]))


def summarize(days: list[DayResult], year: YearData, price: YearPrice,
              mode: str, typ_p: np.ndarray) -> dict[str, float | int]:
    """汇总报告期诊断量（文档 §12.4）：费用分层、价格预报技能、ν 统计、
    星期台阶（性质 9）、日费用形状（性质 10）、紧急购电归因（C4-19）、
    弃电、SOC 与自检。

    Args:
        days: 全部回测日（本函数只取报告期）。
        year: 全年数据（真值与对账自检用）。
        price: 附件4 全年实现电价（价格预报技能与 C4-8 用）。
        mode: 价格口径 r / d / a1。
        typ_p: 附件1 典型日电价（a1 退化口径的结算价，billing_price）。
    """
    rep = report_days(days)
    if not rep:
        raise ValueError("报告期为空（days 须超过暖机期 31 天）")
    checks = [day_checks(d, year) for d in rep]
    u_all = np.array([d.u for d in rep])
    z_day = np.array([d.z for d in rep])
    nu_all = np.array([d.nu_end for d in rep])
    p_hat_err = np.array([np.abs(d.p_hat - price.p[d.j]).mean() for d in rep])
    p_hat_mape = np.array(
        [(np.abs(d.p_hat - price.p[d.j]) / price.p[d.j]).mean() for d in rep])
    emerg_fee = [KAPPA * float(billing_price(price, typ_p, mode, d.j) @ d.u)
                 for d in rep]
    z_online = float(z_day.sum())
    z_ce = float(sum(d.z_ce for d in rep))
    z_expost = float(sum(d.z_expost for d in rep))
    z_pi = float(sum(d.z_pi for d in rep))
    # 紧急购电归因 / 电价分段 / 光伏盈余（§7.6，按紧急电量加权；电价分段
    # 用当日结算价——实现价，a1 退化为附件1 常数价）
    u_energy = float(u_all.sum())
    attr = np.zeros(5)
    bands = np.zeros(len(PRICE_BANDS))
    surplus = 0.0
    for d in rep:
        on = d.u > EPS_EMERG
        if not on.any():
            continue
        net = year.ell[d.j] - year.r[d.j]
        code = attribute_emergency(d.u, d.S, net, d.g)
        for k in range(1, 5):
            attr[k] += d.u[code == k].sum()
        pj = billing_price(price, typ_p, mode, d.j)
        for i, (_, lo, hi) in enumerate(PRICE_BANDS):
            bands[i] += d.u[on & (pj >= lo) & (pj < hi)].sum()
        surplus += d.u[on & (net < 0.0)].sum()
    den = u_energy if u_energy > 0.0 else 1.0
    band_share = bands / den
    # 星期台阶（性质 9）：日末水位按次日星期几分组（周五/周六 = 低价日）
    s144 = {g: [] for g in ("次日周五", "次日周六", "次日周一", "其他")}
    for d in rep:
        nxt_dow = price.dow[d.j + 1] if d.j + 1 < len(price.dow) else 4
        key = {4: "次日周五", 5: "次日周六", 0: "次日周一"}.get(nxt_dow, "其他")
        s144[key].append(d.S[T])
    return {
        "天数": len(rep),
        "Z_online": z_online,
        "Z_ce": z_ce,
        "执行层增益": z_ce - z_online,
        "Z_expost": z_expost,
        "Z_pi": z_pi,
        "EVPI": z_online - z_pi,
        "日前信息损失": z_expost - z_pi,
        "实时信息损失": z_online - z_expost,
        "Z_adj 年末修正": year_end_adjust(days),
        "价格预报 MAE": float(p_hat_err.mean()),
        "价格预报 MAPE": float(p_hat_mape.mean()),
        "ν 均值": float(nu_all.mean()),
        "ν 范围": (float(nu_all.min()), float(nu_all.max())),
        "C4-4 结算价残差 max": max(cash_recheck(d, price, typ_p, mode)
                                   for d in rep),
        "C4-8 ν 区间违例天数": nu_band_check(days, price),
        "C4-14 前缀偏差 max": float(max(c["prefix"] for c in checks)),
        "C6 对账残差 max": float(max(c["c6"] for c in checks)),
        "C7 能量守恒残差 max": float(max(c["c7"] for c in checks)),
        "C4-14 max(c·d)": float(max(c["cd_max"] for c in checks)),
        "C4-14 max(u·w)": float(max(c["uw_max"] for c in checks)),
        "C4-14 u∧c 时段数": int(sum(c["u_and_c"] for c in checks)),
        "C9 序关系违例天数": int(sum(c["c9_bad"] for c in checks)),
        "日费用偏度": float(skew(z_day)),
        "日费用 P95/中位": float(np.quantile(z_day, 0.95) / np.median(z_day)),
        "紧急购电天数": int((u_all > EPS_EMERG).any(axis=1).sum()),
        "紧急购电时段数": int((u_all > EPS_EMERG).sum()),
        "缺口时段频率": float((u_all > EPS_EMERG).sum()) / (len(rep) * T),
        "紧急购电量合计": u_energy,
        "紧急购电费合计": float(sum(emerg_fee)),
        "紧急费占比": float(sum(emerg_fee) / z_online),
        "归因①功率强制占比": float(attr[1] / den),
        "归因②储能耗尽占比": float(attr[2] / den),
        "归因③有电不放占比": float(attr[3] / den),
        "归因④低水位带占比": float(attr[4] / den),
        "紧急购电谷段占比": float(band_share[0]),
        "紧急购电平段占比": float(band_share[1]),
        "紧急购电峰段占比": float(band_share[2]),
        "紧急购电盈余时段占比": float(surplus / den),
        "C4-19 ③为少数": bool(attr[3] / den < 0.5),
        "C4-19 峰段高于谷段": bool(band_share[2] > band_share[0]),
        "C4-19 盈余近零": bool(surplus / den <= 0.05),
        "弃电量合计": float(sum(d.w.sum() for d in rep)),
        "弃电率": float(sum(d.w.sum() for d in rep)
                        / max(float(year.r[[d.j for d in rep]].sum()), 1e-9)),
        "SOC 均值": float(np.mean([d.S for d in rep])),
        "SOC 范围": (float(min(d.S.min() for d in rep)),
                     float(max(d.S.max() for d in rep))),
        "日末 SOC": float(rep[-1].S[T]),
        "星期台阶 次日周五": (float(np.mean(s144["次日周五"]))
                             if s144["次日周五"] else float("nan")),
        "星期台阶 次日周六": (float(np.mean(s144["次日周六"]))
                             if s144["次日周六"] else float("nan")),
        "星期台阶 次日周一": (float(np.mean(s144["次日周一"]))
                             if s144["次日周一"] else float("nan")),
        "星期台阶 其他": float(np.mean(s144["其他"])),
    }


class Result4Xlsx:
    """按官方模板行序填写 result4-2 / result4-3（行号对位不改标签）。

    Attributes:
        template: result4-Q 官方模板路径。
        filename: 输出文件名（result4-2.xlsx / result4-3.xlsx）。
        adjusted: 是否含调整购电量表（仅 4-3）。
    """

    def __init__(self, template: Path, filename: str, adjusted: bool):
        self.template = template
        self.filename = filename
        self.adjusted = adjusted

    def export(self, days: list[DayResult], year: YearData, price: YearPrice,
               out_dir: Path, mode: str = "r",
               typ_p: np.ndarray | None = None) -> Path:
        """写出 result4-Q 并返回路径（费用列一律当日结算价，§9.2 / C4-4）。"""
        rep = report_days(days)
        wb = openpyxl.load_workbook(self.template)
        # 计划表费列 = 名义计划费 Σp^(j)·g（事后结算量，A23）；
        # 调整表费列 = 实际结算购电费 Σ[p·a + ½p|a−g|]（含偏差、不含紧急）
        bill = lambda d: billing_price(price, typ_p, mode, d.j)
        sheets: list[tuple[str, object]] = [
            ("计划购电量", lambda d: (d.g, float(bill(d) @ d.g)))]
        if self.adjusted:
            sheets.append(
                ("调整购电量",
                 lambda d: (d.a, float(bill(d) @ d.a)
                            + 0.5 * float(bill(d) @ np.abs(d.a - d.g)))))
        for name, getter in sheets:
            ws = wb[name]
            for i, day in enumerate(rep):
                row, fee = getter(day)
                for t, val in enumerate(row):
                    ws.cell(row=2 + i, column=2 + t, value=float(val))
                ws.cell(row=2 + i, column=146, value=float(row.sum()))
                ws.cell(row=2 + i, column=147, value=fee)
        write_battery_sheet(wb["充放电量"], rep, year)
        write_emergency_sheet(wb["紧急购电量"], rep, year)
        out_dir.mkdir(parents=True, exist_ok=True)
        path = out_dir / self.filename
        wb.save(path)
        return path


class DetailsCsv:
    """报告期逐日明细 CSV：日期、四层费用、价格诊断、紧急/弃电、SOC。"""

    def export(self, days: list[DayResult], year: YearData, price: YearPrice,
               out_dir: Path, prefix: str = "q4", mode: str = "r",
               typ_p: np.ndarray | None = None) -> Path:
        """写出 q4_details.csv（prefix 区分两线）并返回路径。"""
        rows = []
        for day in report_days(days):
            pj = billing_price(price, typ_p, mode, day.j)
            rows.append({
                "date": str(_py_date(day, year)),
                "z_online": day.z, "z_ce": day.z_ce,
                "z_expost": day.z_expost, "z_pi": day.z_pi,
                "nu_end": day.nu_end,
                "p_mean": float(price.p[day.j].mean()),
                "p_mae": float(np.abs(day.p_hat - price.p[day.j]).mean()),
                "g_sum": float(day.g.sum()),
                **({} if day.a is None else {
                    "a_sum": float(day.a.sum()),
                    "dev_fee": 0.5 * float(pj @ np.abs(day.a - day.g))}),
                "u_sum": float(day.u.sum()),
                "emerg_fee": float(KAPPA * (pj @ day.u)),
                "w_sum": float(day.w.sum()),
                "S0": float(day.S[0]), "S144": float(day.S[T]),
            })
        out_dir.mkdir(parents=True, exist_ok=True)
        path = out_dir / f"{prefix}_details.csv"
        pd.DataFrame(rows).to_csv(path, index=False, float_format="%.6f")
        return path


class SummaryMd:
    """生成题4 结果摘要 markdown（费用分层、价格诊断、自检、紧急归因、
    星期台阶、表 1 四日、口径声明）。"""

    def __init__(self, cfg_label: str, summary: dict, mode: str,
                 line: str):
        self.cfg_label = cfg_label
        self.summary = summary
        self.mode = mode
        self.line = line

    def export(self, days: list[DayResult], year: YearData,
               price: YearPrice, out_dir: Path, prefix: str = "q4") -> Path:
        """写出 q4_summary.md（prefix 区分两线）并返回路径。"""
        s = self.summary
        money = lambda x: f"{x:,.2f} 元"
        pct = lambda x: f"{x * 100:.2f}%"
        core = md_table(["指标", "数值"], [
            ("Z^online（交付，读法 " + self.mode + "）", money(s["Z_online"])),
            ("Z^CE（确定性等价对照）", money(s["Z_ce"])),
            ("执行层增益 Z^CE−Z^online", money(s["执行层增益"])),
            ("Z^ex-post", money(s["Z_expost"])),
            ("Z^PI", money(s["Z_pi"])),
            ("EVPI", money(s["EVPI"])),
            ("Z^adj（年末 SOC 修正，逐日 ν）", money(s["Z_adj 年末修正"])),
            ("日费用偏度 / P95÷中位（性质 10）",
             f"{s['日费用偏度']:.2f} / {s['日费用 P95/中位']:.2f}"),
            ("紧急购电费 / 占比", f"{money(s['紧急购电费合计'])} / "
             f"{pct(s['紧急费占比'])}"),
            ("弃电量 / 弃电率", f"{s['弃电量合计']:,.2f} kWh / "
             f"{pct(s['弃电率'])}"),
            ("SOC 均值 / 日末", f"{s['SOC 均值']:,.1f} / "
             f"{s['日末 SOC']:,.1f} kWh"),
        ])
        price_blk = md_table(["诊断", "数值"], [
            ("价格预报 MAE / MAPE（PR1）",
             f"{s['价格预报 MAE']:.4f} 元/kWh / {pct(s['价格预报 MAPE'])}"),
            ("ν^end 均值 / 范围", f"{s['ν 均值']:.4f} / "
             f"[{s['ν 范围'][0]:.4f}, {s['ν 范围'][1]:.4f}] 元/kWh"),
            ("日末水位 星期台阶（性质 9）",
             f"次日周五 {s['星期台阶 次日周五']:,.0f} / "
             f"次日周六 {s['星期台阶 次日周六']:,.0f} / "
             f"次日周一 {s['星期台阶 次日周一']:,.0f} / "
             f"其他 {s['星期台阶 其他']:,.0f} kWh"),
            ("C4-8 ν 区间违例天数", s["C4-8 ν 区间违例天数"]),
        ])
        checks = md_table(["检查项", "数值"], [
            (k, f"{s[k]:.6f}" if isinstance(s[k], float) else s[k])
            for k in ("C4-4 结算价残差 max", "C4-14 前缀偏差 max",
                      "C6 对账残差 max", "C7 能量守恒残差 max",
                      "C4-14 max(c·d)", "C4-14 max(u·w)",
                      "C4-14 u∧c 时段数", "C9 序关系违例天数")])
        emerg = md_table(["统计量", "数值"], [
            ("发生天数 / 时段数", f"{s['紧急购电天数']} / {s['天数']} 天，"
             f"{s['紧急购电时段数']} 个时段（频率 {pct(s['缺口时段频率'])}）"),
            ("紧急购电量", f"{s['紧急购电量合计']:,.1f} kWh"),
            ("电价分段（按电量，当日实现价）", " / ".join(
                f"{name} {pct(s[key])}" for key, (name, _, _) in zip(
                    ("紧急购电谷段占比", "紧急购电平段占比", "紧急购电峰段占比"),
                    PRICE_BANDS))),
            ("光伏盈余时段占比", pct(s["紧急购电盈余时段占比"])),
        ] + [(ATTR_LABELS[k], pct(s[f"归因{ATTR_LABELS[k]}占比"]))
             for k in range(1, 5)])
        c19 = md_table(["条款（判定性）", "判定", "实测"], [
            ("③有电不放为少数", "过" if s["C4-19 ③为少数"] else "**不过**",
             pct(s["归因③有电不放占比"])),
            ("峰段占比高于谷段", "过" if s["C4-19 峰段高于谷段"] else "**不过**",
             f"峰 {pct(s['紧急购电峰段占比'])} vs "
             f"谷 {pct(s['紧急购电谷段占比'])}"),
            ("盈余时段占比近零", "过" if s["C4-19 盈余近零"] else "**不过**",
             pct(s["紧急购电盈余时段占比"])),
        ])
        tab_rows = []
        for date in PAPER_DATES:
            day = next((d for d in days
                        if str(_py_date(d, year)) == date), None)
            if day is None:
                continue
            supply = day.g if day.a is None else day.a
            tab_rows.append((date,
                             *[f"{day.g[t - 1]:.2f}" for t in PAPER_SLOTS],
                             *([] if day.a is None else
                               [f"{day.a[t - 1]:.2f}" for t in PAPER_SLOTS]),
                             f"{supply.sum():.2f}", f"{day.z:.2f}"))
        heads = (["日期"] + [f"计划 t={t}" for t in PAPER_SLOTS]
                 + ([] if self.line == "4-2" else
                    [f"调整 t={t}" for t in PAPER_SLOTS])
                 + ["全天购电量", "全天购电费"])
        text = f"""# 问题 {self.line} 求解结果（{self.cfg_label}）

## 费用分层（报告期 {s['天数']} 天，价格口径 = {self.mode}）

{core}

## 价格诊断（题4 档 §12.4）

{price_blk}

## 自检（文档 §13）

{checks}

## 紧急购电统计与归因（C4-19）

{emerg}

{c19}

## 论文表 1（四个指定日期；全天购电费为现金费用，实现价结算）

{md_table(heads, tab_rows)}

## 计算口径

- 读法 {self.mode}：{"r = 附件4 实现价按日解锁，0:00 只见预报价（A18 主交付）" if self.mode == "r" else "d = 0:00 已知当日实现价（诊断上界 Z^D）" if self.mode == "d" else "a1 = 附件1 常数价退化（C4-2 复现口径，应逐日复现题" + ("2" if self.line == "4-2" else "3") + "）"}。
- 计划购电费期望按性质 2 折叠为情景均值 p̂ 计价；紧急项逐情景计价（价格
  不确定性只经协方差通道进入目标）。
- 执行层当步用实现价（交易时刻电价）、未来分支用情景价；结算一律实现价
  （C4-4）。
- ν^end 逐日标定（V2 默认：次日预报价循环 LP 对偶，零超参数）；年末
  SOC 修正用 ν^end^(365)，非 0.472（C4-15）。
- 「全天购电费」为事后结算量：0:00 承诺电量 g、单价待实现（A23）；
  计划表费列 = 名义计划费 Σp·g{"" if self.line == "4-2" else "，调整表费列 = 实际结算购电费 Σ[p·a + ½p|a−g|]（含偏差、不含紧急）"}——口径不同，分别声明。
- 充放电量按电池侧填报（ηc、d/η）；紧急购电按极大连续段合并（ε=10⁻⁶ kWh）。
- 暖机期 1 月完整滚动但不计入报告；跨日以真实路径 SOC 链衔接；两线独立滚动（A26）。
"""
        out_dir.mkdir(parents=True, exist_ok=True)
        path = out_dir / f"{prefix}_summary.md"
        path.write_text(text, encoding="utf-8")
        return path
