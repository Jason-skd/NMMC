"""结果导出：统一 Extractor 协议，每个产物一个提取器。

表1/表2 行与诊断量也在此提供——它们是从最优解"提取信息"的公共层，
命令行打印与文件导出共用。
"""
from pathlib import Path
from typing import Protocol

import numpy as np
import openpyxl

from nmmc.data import DayData
from nmmc.lp import ETA, Solution, T

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "docs" / "附件" / "附件5" / "result1.xlsx"

PAPER_SLOTS = [6 * h + 1 for h in (10, 12, 14, 16, 18, 20)]  # t = 61..121


class Extractor(Protocol):
    """从最优解导出一份产物，返回写出的文件路径。"""

    def extract(self, data: DayData, sol: Solution, out_dir: Path) -> Path: ...


def interval_label(t: int) -> str:
    def hm(m):
        return f"{m // 60:02d}:{m % 60:02d}"
    return f"{hm((t - 1) * 10)}-{hm(t * 10)}"


def table1_rows(sol: Solution) -> list[tuple[str, float]]:
    rows = [(interval_label(t), sol.g[t - 1]) for t in PAPER_SLOTS]
    rows += [("全天购电量", sol.g.sum()), ("全天购电费", sol.cost)]
    return rows


def table2_rows(sol: Solution) -> list[tuple[str, float, float]]:
    """6 个 4h 块的 (充电量, 放电量)，电池侧口径（eta*c 与 d/eta）。"""
    rows = []
    for k in range(6):
        seg = slice(24 * k, 24 * (k + 1))
        rows.append((f"{4 * k}:00-{4 * (k + 1)}:00",
                     ETA * sol.c[seg].sum(), sol.d[seg].sum() / ETA))
    return rows


def diagnostics(sol: Solution, data: DayData) -> dict:
    return {
        "Z_base 无电池直买": float(data.p @ np.maximum(data.ell - data.r, 0.0)),
        "弃电量 Σw": float(sol.w.sum()),
        "逐时段 w ≤ r": bool((sol.w <= data.r + 1e-9).all()),
        "充放互斥 max(c·d)": float((sol.c * sol.d).max()),
        "全天损耗": float((1 - ETA) * sol.c.sum() + (1 / ETA - 1) * sol.d.sum()),
        "能量平衡残差": float(
            np.abs(sol.g + data.r + sol.d - data.ell - sol.c - sol.w).max()),
        "SOC 范围": (float(sol.S.min()), float(sol.S.max())),
        "循环误差 |S144-S0|": float(abs(sol.S[T] - sol.S[0])),
    }


class Result1Xlsx:
    """按官方模板行序填写 result1.xlsx（模板行 i ↔ 时段 t=i，不改模板标签）。"""

    def __init__(self, template: Path = TEMPLATE):
        self.template = template

    def extract(self, data: DayData, sol: Solution, out_dir: Path) -> Path:
        wb = openpyxl.load_workbook(self.template)
        ws = wb["计划购电量"]
        for i, val in enumerate(sol.g):
            ws.cell(row=2 + i, column=2, value=float(val))
        ws = wb["充放电量"]
        for k, (_, ch_bat, dis_bat) in enumerate(table2_rows(sol)):
            ws.cell(row=2 + k, column=2, value=float(ch_bat))
            ws.cell(row=2 + k, column=3, value=float(dis_bat))
        ws.cell(row=2, column=5, value=float(sol.S[0]))   # 0:00 储电量
        ws.cell(row=3, column=5, value=float(sol.S[T]))   # 24:00 储电量
        out_dir.mkdir(parents=True, exist_ok=True)
        path = out_dir / "result1.xlsx"
        wb.save(path)
        return path
