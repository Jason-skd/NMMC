"""题2 结算与交付：自检量、紧急购电归因、区间合并、result2 三表、摘要。

交付口径（文档 §8）：result2 计划/真实路径双轨——计划购电量表填 g
（关口侧），充放电量表填 MPC 真实路径按电池侧换算（ηc、d/η，6 个 4h
块 + 0:00/24:00 储电量），紧急购电量表只填 u > ε 的极大连续段。
诊断量与自检清单对应文档 §11，紧急购电归因与电价分段对应 §7.6。
"""
import datetime as dt
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from nmmc.data import ATTACH
from nmmc.data.year import YearData
from nmmc.model.lp import E_MAX_PER, ETA, S_MIN, T
from nmmc.model.q2 import KAPPA
from nmmc.report.common import PAPER_SLOTS, md_table, segment_label
from nmmc.run.q2 import CAP_TOL, DayResult

TEMPLATE = ATTACH / "附件5" / "result2.xlsx"
WARMUP_DAYS = 31      # 1.1–1.31 暖机期，完整模拟但不交付（文档 §1）
EPS_EMERG = 1e-6      # 紧急购电报表阈值，kWh
SOC_RICH = 3000.0     # 归因③"有电不放"的水位判据，kWh（文档 §7.6）
SOC_EPS = 1e-6        # 归因②"储能耗尽"的容差，kWh
PRICE_BANDS = (("谷", 0.0, 0.5), ("平", 0.5, 1.0), ("峰", 1.0, np.inf))
PAPER_DATES = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")
BLOCK_LABELS = [f"{4 * k}:00-{4 * (k + 1)}:00" for k in range(6)]
ATTR_LABELS = {1: "①功率强制", 2: "②储能耗尽", 3: "③a囤电型",
               4: "③b配给型", 5: "④低水位带"}


def report_days(days: list[DayResult]) -> list[DayResult]:
    """报告期切片：j ≥ 31（文档 §1，暖机期不交付）。"""
    return [d for d in days if d.j >= WARMUP_DAYS]


def emergency_segments(u: np.ndarray,
                       eps: float = EPS_EMERG) -> list[tuple[int, int, float]]:
    """紧急购电极大连续段：u_t > ε 的 [i_a, i_b]（0 基含端点）与段内合计。

    Args:
        u: 真实路径紧急购电量，shape (144,)。
        eps: 报表阈值，kWh。

    Returns:
        [(i_a, i_b, Σu)] 列表；无断供日为空。
    """
    on = u > eps
    out = []
    i = 0
    while i < len(u):
        if not on[i]:
            i += 1
            continue
        j = i
        while j + 1 < len(u) and on[j + 1]:
            j += 1
        out.append((i, j, float(u[i:j + 1].sum())))
        i = j + 1
    return out


def attribute_emergency(u: np.ndarray, S: np.ndarray, net: np.ndarray,
                        g: np.ndarray, *,
                        soc_rich: float = SOC_RICH) -> np.ndarray:
    """紧急购电归因（文档 §7.6，必报）：逐时段判定互斥归因码。

    判据按优先级互斥归入：① 放电功率强制 n_t−g_t > D^max（物理下界，
    不可压缩）；② 储能耗尽 S_{t−1} ≤ S̲+ε；③a **囤电型**——电池有电
    （S_{t−1} > soc_rich）且当步 SOC **上升**，即这度电是买来入池的，
    属模型口径缺陷（v6 的 u ≤ (n−g)⁺ 上界令其恒为 0）；③b **配给型**
    ——电池有电但当步 SOC 不升（持平或下降），即存量被留给更贵的窗口，
    是 κ 为平价倍率下的最优跨时段配给，**不得判为失败**；④ 其余
    （低水位带）。0 表示该时段无紧急购电。

    ③a 用"SOC 上升"等价于用 c_t > 0：由 c_td_t = 0，ΔS = ηc − d/η > 0
    当且仅当 c > 0。故 ③a 的时段集恰是性质 8′ 计数的 u∧c 时段集，在
    §5.1 上界下为空——这使 C14a 成为结构性判定条款而非经验阈值。
    ⚠️ "SOC 持平"必须归 ③b 不能归 ③a：上界下 u > 0 ⇒ c = 0 ⇒ SOC 只能
    持平或下降，持平意味着电池把全部存量留给更贵窗口，是配给不是囤电。

    v5 的旧 ③ 把 ③a 与 ③b 焊成一类（判据只看 S_{t−1} > 3000），使
    C14 按字面永不可过；诊断档《题2-SOC与紧急购电诊断》§2.2/§3.3 给出
    两类的机制分离证据（③a 均价 0.71 元/kWh 且 100% 满足套利门槛
    p_t/max p_剩余 < η²；③b 事后 100% 在 p > 1.2 的时段放电）。

    Args:
        u: 真实路径紧急购电量，shape (144,)。
        S: 储电量轨迹，shape (145,)，S[t] 为时段 t 初储电量。
        net: 真实净负荷，shape (144,)。
        g: 计划购电量，shape (144,)。
        soc_rich: ③ 的水位判据，kWh。

    Returns:
        逐时段归因码，shape (144,)，取值 {0,1,2,3,4,5}。
    """
    code = np.zeros(T, dtype=int)
    on = u > EPS_EMERG
    forced = net - g > E_MAX_PER + EPS_EMERG
    empty = S[:-1] <= S_MIN + SOC_EPS
    rich = S[:-1] > soc_rich
    hoard = S[1:] > S[:-1] + SOC_EPS           # 当步 SOC 上升 ⇔ c>0 ⇒ 入池
    base = on & ~forced & ~empty
    code[on & forced] = 1
    code[on & ~forced & empty] = 2
    code[base & rich & hoard] = 3
    code[base & rich & ~hoard] = 4
    code[base & ~rich] = 5
    return code


def day_checks(day: DayResult, year: YearData) -> dict[str, float | int]:
    """单日自检量（文档 §11.3 的 C6/C7/C8/C9）。

    C9 序关系（§7.1）：Z^PI ≤ Z^ex-post ≤ min(Z^online, Z^CE)——
    确定性等价对照同样受事后最优约束。

    Returns:
        dict：C6 对账残差、C7 能量守恒残差、C8 互斥乘积与 u∧c 计数、
        C9 序关系违例标志（1 = 违例）。
    """
    net = year.ell[day.j] - year.r[day.j]
    d_soc = day.S[T] - day.S[0]
    bat = ETA * day.c.sum() - day.d.sum() / ETA
    losses = (1 - ETA) * day.c.sum() + (1 / ETA - 1) * day.d.sum()
    c7 = abs((day.g + day.u - net - day.w).sum() - (d_soc + losses))
    tol = 1e-4
    bad_order = int(day.z_pi > day.z_expost + tol
                    or day.z_expost > day.z + tol
                    or day.z_expost > day.z_ce + tol)
    return {
        "c6": float(abs(bat - d_soc)),
        "c7": float(c7),
        "cd_max": float((day.c * day.d).max()),
        "uw_max": float((day.u * day.w).max()),
        "u_and_c": int(((day.u > EPS_EMERG) & (day.c > EPS_EMERG)).sum()),
        "c9_bad": bad_order,
    }


def summarize(days: list[DayResult], year: YearData, p: np.ndarray) -> dict:
    """汇总报告期诊断量（文档 §11.2）：费用分层、执行层消融、EVPI 分解、
    自检、紧急购电统计与归因、弃电、SOC 水平与预报技能。

    Args:
        days: 全部回测日（本函数只取报告期）。
        year: 全年数据（真值与 SOC 链自检用）。
        p: 逐时段电价（紧急购电费与电价分段用）。
    """
    rep = report_days(days)
    if not rep:
        raise ValueError("报告期为空（days 须超过暖机期 31 天）")
    checks = [day_checks(d, year) for d in rep]
    u_all = np.array([d.u for d in rep])
    emerg_fee = [KAPPA * float(p @ d.u) for d in rep]
    err_l = np.concatenate([np.abs(d.ell_hat - year.ell[d.j]) for d in rep])
    err_r = np.concatenate([np.abs(d.r_hat - year.r[d.j]) for d in rep])
    z_online = float(sum(d.z for d in rep))
    z_ce = float(sum(d.z_ce for d in rep))
    z_expost = float(sum(d.z_expost for d in rep))
    z_pi = float(sum(d.z_pi for d in rep))
    # 紧急购电归因 / 电价分段 / 光伏盈余（§7.6，一律按紧急电量加权）
    u_energy = float(u_all.sum())
    attr = np.zeros(6)
    bands = np.zeros(len(PRICE_BANDS))
    surplus = 0.0
    band_masks = [(p >= lo) & (p < hi) for _, lo, hi in PRICE_BANDS]
    for d in rep:
        on = d.u > EPS_EMERG
        if not on.any():
            continue
        net = year.ell[d.j] - year.r[d.j]
        code = attribute_emergency(d.u, d.S, net, d.g)
        for k in range(1, 6):
            attr[k] += d.u[code == k].sum()
        for i, mask in enumerate(band_masks):
            bands[i] += d.u[on & mask].sum()
        surplus += d.u[on & (net < 0.0)].sum()
    den = u_energy if u_energy > 0.0 else 1.0
    band_share = bands / den
    return {
        "天数": len(rep),
        "Z_online": z_online,
        "Z_ce": z_ce,
        "Z_expost": z_expost,
        "Z_pi": z_pi,
        "执行层增益": z_ce - z_online,
        "EVPI": z_online - z_pi,
        "日前信息损失": z_expost - z_pi,
        "实时信息损失": z_online - z_expost,
        "C6 对账残差 max": float(max(c["c6"] for c in checks)),
        "C7 能量守恒残差 max": float(max(c["c7"] for c in checks)),
        "C8 max(c·d)": float(max(c["cd_max"] for c in checks)),
        "C8 max(u·w)": float(max(c["uw_max"] for c in checks)),
        "性质8 u∧c 时段数": int(sum(c["u_and_c"] for c in checks)),
        "C9 序关系违例天数": int(sum(c["c9_bad"] for c in checks)),
        "C16 上界收敛残差 max": float(max(d.cap_resid for d in rep)),
        "C16 序列线性化轮数 max": int(max(d.cap_iters for d in rep)),
        "C16 未收敛天数": int(sum(d.cap_resid > CAP_TOL for d in rep)),
        "紧急购电天数": int((u_all > EPS_EMERG).any(axis=1).sum()),
        "紧急购电时段数": int((u_all > EPS_EMERG).sum()),
        "缺口时段频率": float((u_all > EPS_EMERG).sum())
        / (len(rep) * T),
        "紧急购电量合计": u_energy,
        "紧急购电费合计": float(sum(emerg_fee)),
        "紧急费占比": float(sum(emerg_fee) / z_online),
        "归因①功率强制占比": float(attr[1] / den),
        "归因②储能耗尽占比": float(attr[2] / den),
        "归因③a囤电型占比": float(attr[3] / den),
        "归因③b配给型占比": float(attr[4] / den),
        "归因④低水位带占比": float(attr[5] / den),
        "紧急购电谷段占比": float(band_share[0]),
        "紧急购电平段占比": float(band_share[1]),
        "紧急购电峰段占比": float(band_share[2]),
        "紧急购电盈余时段占比": float(surplus / den),
        "C14a 囤电型近零": bool(attr[3] / den <= 0.10),
        "C14b 峰段高于谷段": bool(band_share[2] > band_share[0]),
        "C14c 盈余近零": bool(surplus / den <= 0.05),
        "C14d 无紧急购电同时充电": bool(sum(c["u_and_c"] for c in checks) == 0),
        "弃电量合计": float(sum(d.w.sum() for d in rep)),
        "弃电率": float(sum(d.w.sum() for d in rep)
                        / max(float(year.r[[d.j for d in rep]].sum()), 1e-9)),
        "SOC 均值": float(np.mean([d.S for d in rep])),
        "SOC 范围": (float(min(d.S.min() for d in rep)),
                     float(max(d.S.max() for d in rep))),
        "日末 SOC": float(rep[-1].S[T]),
        "负载预报 MAE": float(err_l.mean()),
        "负载预报 p90": float(np.quantile(err_l, 0.9)),
        "光伏预报 MAE": float(err_r.mean()),
        "光伏预报 p90": float(np.quantile(err_r, 0.9)),
    }


def _py_date(day: DayResult, year: YearData) -> dt.date:
    """DayResult 的 0 基日索引转 Python 日期。"""
    return year.dates[day.j].astype(object)


class Result2Xlsx:
    """按官方模板行序填写 result2.xlsx（三表双轨，行号对位不改标签）。

    Attributes:
        template: result2.xlsx 官方模板路径。
    """

    def __init__(self, template: Path = TEMPLATE):
        self.template = template

    def export(self, days: list[DayResult], year: YearData,
               out_dir: Path) -> Path:
        """写出 result2.xlsx（计划购电量 + 充放电量 + 紧急购电量）并返回路径。"""
        rep = report_days(days)
        wb = openpyxl.load_workbook(self.template)
        ws = wb["计划购电量"]
        for i, day in enumerate(rep):
            row = 2 + i
            for t, val in enumerate(day.g):
                ws.cell(row=row, column=2 + t, value=float(val))
            ws.cell(row=row, column=146, value=float(day.g.sum()))
            ws.cell(row=row, column=147, value=float(day.z))
        ws = wb["充放电量"]
        ws.delete_rows(2, ws.max_row - 1)
        for day in rep:
            row = ws.max_row + 1
            for k in range(6):
                seg = slice(24 * k, 24 * (k + 1))
                ws.cell(row=row + k, column=2, value=BLOCK_LABELS[k])
                ws.cell(row=row + k, column=3,
                        value=float(ETA * day.c[seg].sum()))
                ws.cell(row=row + k, column=4,
                        value=float(day.d[seg].sum() / ETA))
            ws.cell(row=row, column=1, value=_py_date(day, year))
            ws.cell(row=row, column=5, value=dt.time(0, 0))
            ws.cell(row=row, column=6, value=float(day.S[0]))
            ws.cell(row=row + 1, column=5, value="24:00")
            ws.cell(row=row + 1, column=6, value=float(day.S[T]))
        ws = wb["紧急购电量"]
        ws.delete_rows(2, ws.max_row - 1)
        for day in rep:
            segs = emergency_segments(day.u)
            row = ws.max_row + 1
            for k, (i_a, i_b, e) in enumerate(segs):
                ws.cell(row=row + k, column=2,
                        value=segment_label(i_a, i_b))
                ws.cell(row=row + k, column=3, value=e)
            if segs:
                ws.cell(row=row, column=1, value=_py_date(day, year))
        out_dir.mkdir(parents=True, exist_ok=True)
        path = out_dir / "result2.xlsx"
        wb.save(path)
        return path


class DetailsCsv:
    """报告期逐日明细 CSV：日期、三层费用、紧急/弃电、SOC、预报误差。"""

    def export(self, days: list[DayResult], year: YearData,
               p: np.ndarray, out_dir: Path) -> Path:
        """写出 q2_details.csv 并返回路径。"""
        rows = []
        for day in report_days(days):
            rows.append({
                "date": str(_py_date(day, year)),
                "z_online": day.z, "z_ce": day.z_ce,
                "z_expost": day.z_expost, "z_pi": day.z_pi,
                "g_sum": float(day.g.sum()),
                "u_sum": float(day.u.sum()),
                "emerg_fee": float(KAPPA * (p @ day.u)),
                "w_sum": float(day.w.sum()),
                "S0": float(day.S[0]), "S144": float(day.S[T]),
                "mae_ell": float(np.abs(day.ell_hat - year.ell[day.j]).mean()),
                "mae_r": float(np.abs(day.r_hat - year.r[day.j]).mean()),
                "cap_resid": day.cap_resid,   # C16，见文档 §6.1
                "cap_iters": day.cap_iters,
            })
        out_dir.mkdir(parents=True, exist_ok=True)
        path = out_dir / "q2_details.csv"
        pd.DataFrame(rows).to_csv(path, index=False, float_format="%.6f")
        return path


class SummaryMd:
    """生成题2 结果摘要 markdown（配置、费用分层与执行层消融、EVPI 分解、
    自检、紧急购电统计与归因、C14 判定、表 1 四日、口径说明）。

    Attributes:
        cfg_label: 配置标签（如 L1-P1-R2-smpc）。
        summary: summarize() 的输出。
    """

    def __init__(self, cfg_label: str, summary: dict):
        self.cfg_label = cfg_label
        self.summary = summary

    def export(self, days: list[DayResult], year: YearData,
               p: np.ndarray, out_dir: Path) -> Path:
        """写出 q2_summary.md 并返回路径。"""
        s = self.summary
        money = lambda x: f"{x:,.2f} 元"
        pct = lambda x: f"{x * 100:.2f}%"
        core = md_table(["指标", "数值"], [
            ("Z^online（交付）", money(s["Z_online"])),
            ("Z^CE（确定性等价对照）", money(s["Z_ce"])),
            ("执行层增益 Z^CE−Z^online", money(s["执行层增益"])),
            ("Z^ex-post", money(s["Z_expost"])),
            ("Z^PI", money(s["Z_pi"])),
            ("EVPI", money(s["EVPI"])),
            ("日前信息损失", money(s["日前信息损失"])),
            ("实时信息损失", money(s["实时信息损失"])),
            ("紧急购电费 / 占比", f"{money(s['紧急购电费合计'])} / "
             f"{pct(s['紧急费占比'])}"),
            ("弃电量 / 弃电率", f"{s['弃电量合计']:,.2f} kWh / "
             f"{pct(s['弃电率'])}"),
            ("SOC 均值 / 范围", f"{s['SOC 均值']:,.1f} kWh / "
             f"{s['SOC 范围'][0]:,.1f} ~ {s['SOC 范围'][1]:,.1f}"),
            ("日末 SOC", f"{s['日末 SOC']:,.1f} kWh"),
        ])
        checks = md_table(["检查项", "数值"], [
            (k, f"{s[k]:.6f}" if isinstance(s[k], float) else s[k])
            for k in ("C6 对账残差 max", "C7 能量守恒残差 max", "C8 max(c·d)",
                      "C8 max(u·w)", "性质8 u∧c 时段数", "C9 序关系违例天数",
                      "C16 上界收敛残差 max", "C16 序列线性化轮数 max",
                      "C16 未收敛天数")])
        emerg = md_table(["统计量", "数值"], [
            ("发生天数 / 时段数", f"{s['紧急购电天数']} / {s['天数']} 天，"
             f"{s['紧急购电时段数']} 个时段（频率 {pct(s['缺口时段频率'])}）"),
            ("紧急购电量", f"{s['紧急购电量合计']:,.1f} kWh"),
            ("电价分段（按电量）", " / ".join(
                f"{name} {pct(s[key])}" for key, (name, _, _) in zip(
                    ("紧急购电谷段占比", "紧急购电平段占比", "紧急购电峰段占比"),
                    PRICE_BANDS))),
            ("光伏盈余时段占比", pct(s["紧急购电盈余时段占比"])),
        ] + [(ATTR_LABELS[k], pct(s[f"归因{ATTR_LABELS[k]}占比"]))
             for k in range(1, 6)])
        c14 = md_table(["条款（判定性）", "判定", "实测"], [
            ("③a 囤电型近零（≤10%）",
             "过" if s["C14a 囤电型近零"] else "**不过**",
             pct(s["归因③a囤电型占比"])),
            ("峰段占比高于谷段", "过" if s["C14b 峰段高于谷段"] else "**不过**",
             f"峰 {pct(s['紧急购电峰段占比'])} vs "
             f"谷 {pct(s['紧急购电谷段占比'])}"),
            ("盈余时段占比近零", "过" if s["C14c 盈余近零"] else "**不过**",
             pct(s["紧急购电盈余时段占比"])),
            ("无紧急购电同时充电（性质 8′）",
             "过" if s["C14d 无紧急购电同时充电"] else "**不过**",
             f"{s['性质8 u∧c 时段数']} 个时段"),
        ])
        skill = md_table(["预报", "MAE (kWh/时段)", "p90 (kWh/时段)"], [
            ("负载", f"{s['负载预报 MAE']:.2f}", f"{s['负载预报 p90']:.2f}"),
            ("光伏", f"{s['光伏预报 MAE']:.2f}", f"{s['光伏预报 p90']:.2f}")])
        tab_rows = []
        for date in PAPER_DATES:
            day = next((d for d in days
                        if str(_py_date(d, year)) == date), None)
            if day is None:
                continue
            vals = [f"{day.g[t - 1]:.2f}" for t in PAPER_SLOTS]
            tab_rows.append((date, *vals, f"{day.g.sum():.2f}",
                             f"{day.z:.2f}"))
        table1 = md_table(
            ["日期"] + [f"t={t}" for t in PAPER_SLOTS] + ["全天购电量", "购电费"],
            tab_rows)
        text = f"""# 问题 2 求解结果（{self.cfg_label}）

## 费用分层（报告期 {s['天数']} 天）

{core}

## 自检（文档 §11.3）

{checks}

## 紧急购电统计与归因（文档 §7.6）

{emerg}

## C14 归因健康度（判定性，任一违反不得交付）

{c14}

## 预报技能

{skill}

## 论文表 1（四个指定日期，计划口径；购电费含紧急费）

{table1}

## 计算口径

- Z = Σp·g + κΣp·u（现金费用，真实路径 u）；Z̃ 含末端估值虚拟记账项，不上报。
- Z^CE 与 Z^online 同 g、同日初 SOC，只差未来表示（单条 α^rt 分位路径 vs
  情景分布，文档 §7.3），费用差为执行层机制的干净归因。
- 紧急购电口径：u_t ≤ (n_t−g_t)⁺（文档 §5.1，假设 A6′）——四条臂
  （交付 / CE / 事后 / PI）同口径；③b 配给型只报告不判定（§7.6）。
- 归因与电价分段一律按紧急电量加权；"近零"按 ≤5%（③a 按 ≤10%）执行。
- 充放电量按电池侧填报（ηc、d/η）；紧急购电按极大连续段合并（ε=10⁻⁶ kWh）。
- 暖机期 1 月完整滚动但不计入报告；跨日以真实路径 SOC 链衔接。
- result2.xlsx 按官方模板行号对位填写，不改动模板标签。
"""
        out_dir.mkdir(parents=True, exist_ok=True)
        path = out_dir / "q2_summary.md"
        path.write_text(text, encoding="utf-8")
        return path
